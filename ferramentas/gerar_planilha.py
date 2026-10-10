import json, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

D = json.load(open('carga.json'))
wb = Workbook()
AZUL = Font(color='0000FF')
CAB = Font(bold=True, color='FFFFFF')
FCAB = PatternFill('solid', fgColor='0F5C8C')
FIN = PatternFill('solid', fgColor='FFF9E6')
fino = Side(style='thin', color='DFE4EA')
NMAX = 400      # linhas preparadas em Resultados
AV = 6000       # linhas de referencia em Avaliacoes

def cab(ws, headers, row=1, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h); c.font = CAB; c.fill = FCAB
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = ws.cell(row + 1, 1)
    for i, w in enumerate(widths or [], 1): ws.column_dimensions[L(i)].width = w

def data(s): return dt.date.fromisoformat(s) if s else None

# ---------------------------------------------------------------- LEIA-ME
ws = wb.active; ws.title = 'LEIA-ME'
texto = [
    ('CONVISITA | Base de dados das visitas do COMSAÚDE Americana', True),
    ('Versão em planilha do mesmo modelo de dados do sistema web (pasta convisita). Cada informação é registrada uma única vez e os resultados são calculados por fórmula.', False),
    ('', False),
    ('Abas de cadastro e lançamento (texto azul = dado digitado):', True),
    ('Dimensoes: as quatro dimensões do instrumento, com peso (soma = 100%) e pontuação máxima por item.', False),
    ('Itens: os 14 itens indicadores (I.1 a IV.2), com meta esperada e fonte de observação.', False),
    ('Faixas: limites do farol (vermelho, amarelo, verde) usados na classificação.', False),
    ('Unidades: cadastro único dos estabelecimentos, com tipo e periodicidade de visita em meses.', False),
    ('Visitas: uma linha por visita (id, unidade, ciclo, data, situação).', False),
    ('Avaliacoes: uma linha por item avaliado em cada visita (14 linhas por visita). Pontuação 0, 1 ou 2; NA = 1 quando o item não se aplica; pontuação vazia com NA = 0 significa item não avaliado.', False),
    ('', False),
    ('Abas calculadas (não digitar):', True),
    ('Resultados: pontos, percentual por dimensão e nota final de cada visita. A nota final é a média ponderada das dimensões avaliadas; quando uma dimensão inteira não se aplica, seu peso é redistribuído entre as demais.', False),
    ('Painel: última visita de cada unidade nos ciclos indicados nas células C3 e F3, variação entre ciclos, próxima visita prevista e situação da agenda.', False),
    ('Origem_Abas: rastreabilidade das 43 abas da planilha original (importada ou não, e por quê).', False),
    ('', False),
    ('Como lançar uma nova visita:', True),
    ('1. Em Visitas, acrescente uma linha com novo id (sequencial), nome da unidade exatamente como em Unidades, ciclo, data e situação (concluida, incompleta, rascunho).', False),
    ('2. Em Avaliacoes, acrescente as 14 linhas da visita (mesmo id), uma por item, com a dimensão (I, II, III ou IV) e a pontuação.', False),
    ('3. Resultados e Painel se atualizam automaticamente (Resultados está preparado para até 400 visitas).', False),
]
for i, (t, b) in enumerate(texto, 1):
    c = ws.cell(i, 1, t); c.font = Font(bold=b, size=13 if i == 1 else 11); c.alignment = Alignment(wrap_text=True)
ws.column_dimensions['A'].width = 130

# ---------------------------------------------------------------- Dimensoes
ws = wb.create_sheet('Dimensoes')
cab(ws, ['Código', 'Dimensão', 'Peso', 'Nº itens', 'Pontuação máx. por item', 'Pontos máx. da dimensão'], widths=[9, 52, 10, 10, 14, 14])
for r, d in enumerate(D['dims'], 2):
    ws.cell(r, 1, d[1]); ws.cell(r, 2, d[2]); c = ws.cell(r, 3, d[4]); c.font = AZUL; c.number_format = '0%'
    ws.cell(r, 4, f'=COUNTIF(Itens!$B$2:$B$100,A{r})')
    c = ws.cell(r, 5, 2); c.font = AZUL
    ws.cell(r, 6, f'=D{r}*E{r}')
ws['B6'] = 'Total'; ws['B6'].font = Font(bold=True)
ws['C6'] = '=SUM(C2:C5)'; ws['C6'].number_format = '0%'
ws['D6'] = '=SUM(D2:D5)'; ws['F6'] = '=SUM(F2:F5)'

# ---------------------------------------------------------------- Itens
ws = wb.create_sheet('Itens')
cab(ws, ['Item', 'Dimensão', 'Nome do item indicador', 'Meta esperada', 'Fonte de observação', 'Periodicidade'], widths=[8, 10, 34, 60, 50, 30])
dcod = {1: 'I', 2: 'II', 3: 'III', 4: 'IV'}
for r, it in enumerate(D['itens'], 2):
    for c, v in enumerate([it['codigo'], dcod[it['dim']], it['nome'], it['meta'], it['fonte'], it['periodicidade']], 1):
        x = ws.cell(r, c, v); x.alignment = Alignment(wrap_text=True, vertical='top')

# ---------------------------------------------------------------- Faixas
ws = wb.create_sheet('Faixas')
cab(ws, ['Cor', 'Rótulo', 'Limite inferior (inclusive)'], widths=[12, 30, 22])
for r, (cor, rot, li) in enumerate([('vermelho', 'Crítico', 0), ('amarelo', 'Atenção', 0.4), ('verde', 'Adequado', 0.8)], 2):
    ws.cell(r, 1, cor); ws.cell(r, 2, rot); c = ws.cell(r, 3, li); c.font = AZUL; c.number_format = '0%'
ws['A6'] = 'A planilha original usava formatação condicional com intervalos 0 a 0,40 / 0,41 a 0,79 / 0,80 a 1, que deixavam lacunas sem cor (ex.: 79,5%). Aqui os limites são contínuos.'

# ---------------------------------------------------------------- Unidades
ws = wb.create_sheet('Unidades')
cab(ws, ['id', 'Unidade', 'Tipo', 'Aba na planilha original', 'Situação', 'Periodicidade (meses)', 'Gestão', 'Observação'], widths=[5, 48, 34, 24, 10, 12, 14, 70])
for r, (n, t, a, o) in enumerate(D['unidades'], 2):
    for c, v in enumerate([r - 1, n, t, (a or '').strip() or None, 'ativa', 6, None, o], 1):
        x = ws.cell(r, c, v)
        if c in (5, 6, 7): x.font = AZUL
NU = len(D['unidades']) + 1
dvu = DataValidation(type='list', formula1=f'=Unidades!$B$2:$B${NU + 60}', allow_blank=False)

# ---------------------------------------------------------------- Visitas
ws = wb.create_sheet('Visitas')
cab(ws, ['id', 'Unidade', 'Ciclo', 'Data da visita', 'Data aproximada', 'Situação', 'Origem', 'Observação da importação'], widths=[6, 46, 8, 13, 10, 12, 44, 90])
for r, v in enumerate(D['visitas'], 2):
    vals = [v['id'], v['unidade'], str(v['ciclo']), data(v['data']), 'sim' if v['aprox'] else 'não', v['situacao'], v['origem'], v['obs']]
    for c, x in enumerate(vals, 1):
        cell = ws.cell(r, c, x); cell.font = AZUL
        if c == 4: cell.number_format = 'DD/MM/YYYY'
ws.add_data_validation(dvu); dvu.add(f'B2:B{NMAX + 1}')
dvs = DataValidation(type='list', formula1='"concluida,incompleta,rascunho,em_conferencia"'); ws.add_data_validation(dvs); dvs.add(f'F2:F{NMAX + 1}')

# ---------------------------------------------------------------- Avaliacoes
ws = wb.create_sheet('Avaliacoes')
cab(ws, ['visita_id', 'Unidade (conferência)', 'Item', 'Nome do item', 'Ciclo (conferência)', 'Dimensão', 'Pontuação (0 a 2)', 'Não se aplica (1/0)'], widths=[9, 40, 7, 34, 10, 10, 12, 12])
nomes = {it['codigo']: it['nome'] for it in D['itens']}
r = 2
for v in D['visitas']:
    for it in D['itens']:
        x = v['itens'][it['codigo']]
        ws.cell(r, 1, v['id']).font = AZUL
        ws.cell(r, 2, f'=IFERROR(INDEX(Visitas!$B$2:$B${NMAX + 1},MATCH(A{r},Visitas!$A$2:$A${NMAX + 1},0)),"")')
        ws.cell(r, 3, it['codigo']).font = AZUL
        ws.cell(r, 4, f'=IFERROR(INDEX(Itens!$C$2:$C$100,MATCH(C{r},Itens!$A$2:$A$100,0)),"")')
        ws.cell(r, 5, f'=IFERROR(INDEX(Visitas!$C$2:$C${NMAX + 1},MATCH(A{r},Visitas!$A$2:$A${NMAX + 1},0)),"")')
        ws.cell(r, 6, f'=IFERROR(INDEX(Itens!$B$2:$B$100,MATCH(C{r},Itens!$A$2:$A$100,0)),"")')
        c = ws.cell(r, 7, x['p']); c.font = AZUL
        c = ws.cell(r, 8, 1 if x['na'] else 0); c.font = AZUL
        r += 1
dvp = DataValidation(type='whole', operator='between', formula1='0', formula2='2', allow_blank=True); ws.add_data_validation(dvp); dvp.add(f'G2:G{AV}')
ws.auto_filter.ref = f'A1:H{r - 1}'

# ---------------------------------------------------------------- Resultados
ws = wb.create_sheet('Resultados')
heads = ['visita_id', 'Ciclo', 'Data', 'Unidade', 'Situação',
         'Pontos I', 'Pontos II', 'Pontos III', 'Pontos IV',
         'Possíveis I', 'Possíveis II', 'Possíveis III', 'Possíveis IV',
         '% Infra. Física', '% Infra. Humana', '% Plano Municipal', '% Percepção Usuário',
         'Nota final', 'Classificação', 'Nota pelo método da planilha original', 'Válida p/ painel']
cab(ws, heads, widths=[8, 7, 11, 40, 11] + [8] * 8 + [11] * 4 + [10, 13, 15, 9])
A = f'Avaliacoes!$A$2:$A${AV}'; F = f'Avaliacoes!$F$2:$F${AV}'; G = f'Avaliacoes!$G$2:$G${AV}'; H = f'Avaliacoes!$H$2:$H${AV}'
for r in range(2, NMAX + 2):
    vr = r
    ws.cell(r, 1, f'=IF(Visitas!A{vr}="","",Visitas!A{vr})')
    ws.cell(r, 2, f'=IF($A{r}="","",Visitas!C{vr})')
    c = ws.cell(r, 3, f'=IF($A{r}="","",Visitas!D{vr})'); c.number_format = 'DD/MM/YYYY'
    ws.cell(r, 4, f'=IF($A{r}="","",Visitas!B{vr})')
    ws.cell(r, 5, f'=IF($A{r}="","",Visitas!F{vr})')
    for k, cod in enumerate(['I', 'II', 'III', 'IV']):
        dr = k + 2
        ws.cell(r, 6 + k, f'=IF($A{r}="","",SUMIFS({G},{A},$A{r},{F},"{cod}",{H},0))')
        ws.cell(r, 10 + k, f'=IF($A{r}="","",COUNTIFS({A},$A{r},{F},"{cod}",{H},0,{G},"<>")*Dimensoes!$E${dr})')
        c = ws.cell(r, 14 + k, f'=IF($A{r}="","",IF({L(10 + k)}{r}>0,{L(6 + k)}{r}/{L(10 + k)}{r},""))'); c.number_format = '0.0%'
    num = '+'.join(f'IF({L(14 + k)}{r}="",0,{L(14 + k)}{r}*Dimensoes!$C${k + 2})' for k in range(4))
    den = '+'.join(f'IF({L(14 + k)}{r}="",0,Dimensoes!$C${k + 2})' for k in range(4))
    c = ws.cell(r, 18, f'=IF($A{r}="","",IF(({den})=0,"",({num})/({den})))'); c.number_format = '0.0%'; c.font = Font(bold=True)
    ws.cell(r, 19, f'=IF(R{r}="","",IF(R{r}<Faixas!$C$3,"vermelho",IF(R{r}<Faixas!$C$4,"amarelo","verde")))')
    c = ws.cell(r, 20, f'=IF($A{r}="","",{num})'); c.number_format = '0.0%'
    ws.cell(r, 21, f'=IF($A{r}="","",IF(OR(E{r}="concluida",E{r}="incompleta"),1,0))')

def farois(ws, rng, col_ref):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col_ref}),{col_ref}<Faixas!$C$3)'], fill=PatternFill('solid', fgColor='FBE4E1'), font=Font(color='C0392B')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col_ref}),{col_ref}>=Faixas!$C$3,{col_ref}<Faixas!$C$4)'], fill=PatternFill('solid', fgColor='FFF4D1'), font=Font(color='8A6500')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col_ref}),{col_ref}>=Faixas!$C$4)'], fill=PatternFill('solid', fgColor='E3F3EA'), font=Font(color='2E8B57')))
farois(ws, f'N2:R{NMAX + 1}', 'N2')
ws.auto_filter.ref = f'A1:U{NMAX + 1}'

# ---------------------------------------------------------------- Painel
ws = wb.create_sheet('Painel', 1)
ws['A1'] = 'Painel de desempenho por unidade'; ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = 'Última visita válida de cada unidade nos ciclos escolhidos (altere C3 e F3). Farol conforme aba Faixas.'
ws['B3'] = 'Ciclo base:'; ws['C3'] = '2023'; ws['C3'].font = AZUL; ws['C3'].fill = FIN
ws['E3'] = 'Ciclo comparado:'; ws['F3'] = '2026'; ws['F3'].font = AZUL; ws['F3'].fill = FIN
ws['J3'] = 'Data de referência:'; ws['K3'] = '=TODAY()'; ws['K3'].number_format = 'DD/MM/YYYY'
heads = ['Unidade', 'Tipo', 'Data (base)', 'Nota (base)', 'Data (comparado)', 'Nota (comparado)', 'Variação (p.p.)',
         'Última visita', 'Próxima prevista', 'Situação da agenda',
         '% Infra. Física', '% Infra. Humana', '% Plano Municipal', '% Percepção Usuário']
cab(ws, heads, row=5, widths=[46, 30, 11, 10, 11, 10, 10, 12, 12, 16, 11, 11, 11, 11])
R = f'Resultados!$A$2:$A${NMAX + 1}'
RB = f'Resultados!$B$2:$B${NMAX + 1}'; RC = f'Resultados!$C$2:$C${NMAX + 1}'; RD = f'Resultados!$D$2:$D${NMAX + 1}'
RR = f'Resultados!$R$2:$R${NMAX + 1}'; RU = f'Resultados!$U$2:$U${NMAX + 1}'
for i in range(len(D['unidades'])):
    r = 6 + i; ur = i + 2
    ws.cell(r, 1, f'=Unidades!B{ur}'); ws.cell(r, 2, f'=Unidades!C{ur}')
    for (dc, nc, cic) in [(3, 4, '$C$3'), (5, 6, '$F$3')]:
        c = ws.cell(r, dc, f'=IF(_xlfn.MAXIFS({RC},{RD},$A{r},{RB},{cic},{RU},1)=0,"",_xlfn.MAXIFS({RC},{RD},$A{r},{RB},{cic},{RU},1))'); c.number_format = 'DD/MM/YYYY'
        c = ws.cell(r, nc, f'=IF({L(dc)}{r}="","",SUMIFS({RR},{RD},$A{r},{RB},{cic},{RC},{L(dc)}{r},{RU},1))'); c.number_format = '0.0%'
    c = ws.cell(r, 7, f'=IF(OR(D{r}="",F{r}=""),"",(F{r}-D{r})*100)'); c.number_format = '+0.0;-0.0;0.0'
    c = ws.cell(r, 8, f'=IF(_xlfn.MAXIFS({RC},{RD},$A{r},{RU},1)=0,"",_xlfn.MAXIFS({RC},{RD},$A{r},{RU},1))'); c.number_format = 'DD/MM/YYYY'
    c = ws.cell(r, 9, f'=IF(H{r}="","",EDATE(H{r},Unidades!F{ur}))'); c.number_format = 'DD/MM/YYYY'
    ws.cell(r, 10, f'=IF(Unidades!E{ur}<>"ativa","suspensa",IF(H{r}="","sem visita",IF(I{r}<$K$3,"atrasada",IF(I{r}<=$K$3+30,"próximos 30 dias","em dia"))))')
    for k, col in enumerate(['N', 'O', 'P', 'Q']):
        c = ws.cell(r, 11 + k, f'=IF(E{r}="","",SUMIFS(Resultados!${col}$2:${col}${NMAX + 1},{RD},$A{r},{RB},$F$3,{RC},E{r},{RU},1))'); c.number_format = '0.0%'
last = 5 + len(D['unidades'])
farois(ws, f'D6:D{last}', 'D6'); farois(ws, f'F6:F{last}', 'F6'); farois(ws, f'K6:N{last}', 'K6')
ws.conditional_formatting.add(f'J6:J{last}', FormulaRule(formula=['J6="atrasada"'], fill=PatternFill('solid', fgColor='FBE4E1')))
ws.conditional_formatting.add(f'J6:J{last}', FormulaRule(formula=['J6="próximos 30 dias"'], fill=PatternFill('solid', fgColor='FFF4D1')))
ws.auto_filter.ref = f'A5:N{last}'
r = last + 2
ws.cell(r, 1, 'Síntese do ciclo comparado').font = Font(bold=True)
ws.cell(r + 1, 1, 'Unidades visitadas'); ws.cell(r + 1, 4, f'=COUNT(F6:F{last})')
ws.cell(r + 2, 1, 'Nota média'); c = ws.cell(r + 2, 4, f'=IFERROR(AVERAGE(F6:F{last}),"")'); c.number_format = '0.0%'
ws.cell(r + 3, 1, 'Adequadas (verde)'); ws.cell(r + 3, 4, f'=COUNTIF(F6:F{last},">="&Faixas!$C$4)')
ws.cell(r + 4, 1, 'Atenção (amarelo)'); ws.cell(r + 4, 4, f'=COUNTIFS(F6:F{last},">="&Faixas!$C$3,F6:F{last},"<"&Faixas!$C$4)')
ws.cell(r + 5, 1, 'Críticas (vermelho)'); ws.cell(r + 5, 4, f'=COUNTIF(F6:F{last},"<"&Faixas!$C$3)')
ws.cell(r + 6, 1, 'Visitas atrasadas'); ws.cell(r + 6, 4, f'=COUNTIF(J6:J{last},"atrasada")')

# ---------------------------------------------------------------- Origem_Abas
ws = wb.create_sheet('Origem_Abas')
cab(ws, ['Nº', 'Aba original', 'Unidade no cadastro', 'Nome no quadro-resumo da aba', 'Data no rodapé', 'Notas I.1 a IV.2 (N = não se aplica)', 'Tratamento', 'Motivo'], widths=[5, 24, 40, 26, 12, 22, 14, 90])
excl = {e['aba']: e['motivo'] for e in D['excluidas']}
amap = {a.strip(): n for n, t, a, o in D['unidades'] if a}
for i, a in enumerate(D['abas'], 1):
    r = i + 1
    vals = [i, a['aba'], amap.get(a['aba'].strip()), a['cabecalho'], data(a['data']), a['vetor'],
            'não importada' if a['aba'] in excl else 'importada', excl.get(a['aba'])]
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, v)
        if c == 5: x.number_format = 'DD/MM/YYYY'

out = '../planilha/CONVISITA_base_visitas.xlsx'
wb.save(out); print(out, wb.sheetnames)
