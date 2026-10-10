"""Le a planilha original e gera os scripts SQL de referencia e carga historica,
alem de um JSON intermediario usado para montar a planilha normalizada."""
import openpyxl, re, json, datetime as dt

SRC = '../planilha/original/Pontuação visitas 2026atualizada.xlsx'  # executar a partir da pasta ferramentas
SQLDIR = '../sql/'
wb = openpyxl.load_workbook(SRC)

def q(v):
    if v is None: return 'NULL'
    if isinstance(v, bool): return '1' if v else '0'
    if isinstance(v, (int, float)): return str(v)
    if isinstance(v, (dt.date, dt.datetime)): return "'" + v.strftime('%Y-%m-%d') + "'"
    s = str(v).strip().replace('\\', '\\\\').replace("'", "''")
    s = re.sub(r'\s+\n', '\n', s)
    return "'" + s + "'"

# ------------------------------------------------------------------ instrumento
ws = wb['UBS Vila Gallo']  # aba com todos os nomes de item preenchidos
dims = [
    (1, 'I', 'Infraestrutura Física (prédios, equipamentos e insumos)', 'Infraestrutura Física', 0.25, 1),
    (2, 'II', 'Infraestrutura Humana', 'Infraestrutura Humana', 0.25, 2),
    (3, 'III', 'Desempenho no Plano Municipal de Saúde', 'Plano Municipal de Saúde', 0.25, 3),
    (4, 'IV', 'Percepção do Usuário', 'Percepção do Usuário', 0.25, 4),
]
dim_of = lambda code: {'I': 1, 'II': 2, 'III': 3, 'IV': 4}[code.split('.')[0].strip()]
itens = []
for k, r in enumerate(range(4, 18), start=1):
    code = re.sub(r'\s+', '', ws.cell(r, 2).value).rstrip('.')
    itens.append(dict(id=k, codigo=code, dim=dim_of(code), nome=ws.cell(r, 3).value.strip(),
                      meta=ws.cell(r, 4).value.strip(), fonte=ws.cell(r, 5).value.strip(),
                      periodicidade=ws.cell(r, 6).value.strip(), ordem=k))
ITEM_CODES = [i['codigo'] for i in itens]

tipos = ['Unidade Básica de Saúde (UBS)', 'Estratégia Saúde da Família (ESF)', 'Centro de Atenção Psicossocial (CAPS)',
         'Urgência e Emergência', 'Hospital', 'Atenção Especializada', 'Vigilância em Saúde',
         'Apoio Diagnóstico e Terapêutico', 'Assistência Farmacêutica', 'Atenção Domiciliar',
         'Entidade Parceira / Conveniada', 'A classificar']
T = {n: i + 1 for i, n in enumerate(tipos)}
UBS, ESF, CAPS, URG, HOSP, ESP, VIG, APOIO, FARM, DOM, ENT, OUT = tipos

# ------------------------------------------------------------------ unidades
# (nome canonico, tipo, aba da planilha, observacao)
unidades = [
    ('UAD - Unidade de Atendimento Domiciliar', DOM, 'UAD', None),
    ('UBS Jardim São Paulo', UBS, 'UBS J. São Paulo', None),
    ('Hospital Municipal Dr. Waldemar Tebaldi', HOSP, 'H. M.Urg. e Emerg.', 'Cronograma de visitas próprio (Performance 2023, linha 41).'),
    ('UPA Dona Rosa', URG, 'UPA D. Rosa', 'Verificar se corresponde à UPA São José avaliada em 2023.'),
    ('Banco de Sangue', APOIO, 'Banco de Sangue', None),
    ('UBS Parque Gramado', UBS, 'UBS Parque Gramado', 'Situação em 2023: reforma.'),
    ('ESF São José', ESF, 'ESF São José', 'Unidade recém reformada e entregue em novembro/2023.'),
    ('ESF Praia Azul', ESF, 'ESF Praia Azul', None),
    ('ESF Jardim Brasil', ESF, 'ESF Jardim Brasil', None),
    ('CAPS Adulto', CAPS, 'Caps Adulto', 'Verificar se corresponde ao CAPS Arte & Vida da lista 2023.'),
    ('UBS Mathiensen', UBS, 'UBS Mathiensen', None),
    ('UBS São Vito', UBS, 'UBS São Vito', 'Situação em 2023: reforma.'),
    ('CAPS Infantil Larissa de Almeida', CAPS, 'Caps Infantil', None),
    ('UBS Parque das Nações', UBS, 'UBS Parque das Nações', 'Situação em 2023: reforma.'),
    ('ESF Liberdade', ESF, 'ESF Liberdade', 'Situação em 2023: reforma.'),
    ('ESF Jaguari', ESF, 'ESF Jaguari', 'Situação em 2023: reforma.'),
    ('ESF Zanaga I', ESF, 'ESF Zanaga I', None),
    ('ESF Zanaga II', ESF, 'ESF Zanaga II', None),
    ('ADAM', ENT, 'ADAM', None),
    ('AEPHIVA', ENT, 'AEPHIVA', 'Visita extraordinária registrada a partir de ago/2023.'),
    ('UBS Cariobinha', UBS, 'UBS Cariobinha ', 'Unidade recém reformada e entregue em dezembro/2023.'),
    ('UNACON - Unidade de Oncologia', ESP, 'UNACON', None),
    ('APAE', ENT, 'APAE', 'Visita extraordinária registrada a partir de ago/2023.'),
    ('UVISA - Vigilância Sanitária', VIG, 'UVISA', None),
    ('SAE/DST', ESP, 'SAE DST', 'Itens de percepção do usuário não se aplicam (sigilo).'),
    ('Clínica Jones', OUT, 'Clínica Jones', None),
    ('UBS Vila Gallo', UBS, 'UBS Vila Gallo', None),
    ('CCZ - Centro de Controle de Zoonoses', VIG, 'CCZ ', None),
    ('CTA - Centro de Testagem e Aconselhamento', ESP, 'CTA', None),
    ('UBS Jardim Ipiranga', UBS, 'UBS Jardim Ipiranga', None),
    ('CAPS AD Vida Nova (Álcool e Drogas)', CAPS, 'CAPS Alcool e Drogas', None),
    ('PA Antonio Zanaga', URG, 'PA Zanaga', 'Denominada UPA Bairro Antonio Zanaga em 2023.'),
    ('UBS Dona Rosa', UBS, 'UBS Dona Rosa', None),
    ('UBS Jardim Boer', UBS, 'UBS Jd Boer ', None),
    ('ESF Alvorada', ESF, 'ESF Alvorada ', None),
    ('ESF Mário Covas', ESF, 'ESF Mário Covas', None),
    ('Núcleo de Especialidades de Americana', ESP, 'Núcleo Especialidades ', None),
    ('UBS São Domingos', UBS, 'UBS São Domingos ', None),
    ('UBS Guanabara', UBS, 'UBS Guanabara ', None),
    ('UBS Liberdade', UBS, 'UBS LIBERDADE', None),
    ('Clínica Modelo Zanaga', OUT, 'Clínica Mod.  ZANAGA', 'Confirmar nome oficial e tipo de estabelecimento.'),
    ('Farmácia Central', FARM, 'Farmácia Central (2)', 'Situação em 2023: reforma.'),
    ('CEO - Centro de Especialidades Odontológicas', ESP, 'CEO', 'Situação em 2023: reforma.'),
    ('UPA São José', URG, None, 'Avaliada em 2023. Sem aba própria na planilha de 2026.'),
    ('SAMU 192 - Serviço de Atendimento Móvel de Urgência', URG, None, 'Constava na lista 2023 sem visita (situação: reforma).'),
    ('Serviço de Transporte em Saúde', APOIO, None, 'Visita extraordinária registrada a partir de ago/2023.'),
]
U = {n: i + 1 for i, (n, *_) in enumerate(unidades)}
ABA2U = {a: n for n, _, a, _ in unidades if a}

# ------------------------------------------------------------------ visitas 2023
ag = wb['Avaliação Global']
ag_alias = {'UAD': 'UAD - Unidade de Atendimento Domiciliar', 'UBS Mathiensen': 'UBS Mathiensen',
            'UBS Cariobinha': 'UBS Cariobinha', 'UNACON': 'UNACON - Unidade de Oncologia',
            'UPA São José': 'UPA São José', 'UBS Vila Gallo': 'UBS Vila Gallo',
            'Centro de Zoonoses': 'CCZ - Centro de Controle de Zoonoses',
            'CTA': 'CTA - Centro de Testagem e Aconselhamento', 'UBS Jd Ipiranga': 'UBS Jardim Ipiranga',
            'CAPS Alcool e drogas': 'CAPS AD Vida Nova (Álcool e Drogas)', 'PA Zanaga': 'PA Antonio Zanaga',
            'UBS Dona Rosa': 'UBS Dona Rosa', 'UBS Jd Boer': 'UBS Jardim Boer', 'ESF Alvorada': 'ESF Alvorada',
            'ESF Mário Covas': 'ESF Mário Covas', 'Núcleo Especialidades': 'Núcleo de Especialidades de Americana',
            'UBS São Domingos': 'UBS São Domingos', 'UBS Guanabara': 'UBS Guanabara',
            'UBS Jd. São Paulo': 'UBS Jardim São Paulo'}
ag_scores = {}
for r in range(2, ag.max_row + 1):
    a = ag.cell(r, 1).value
    if not a: continue
    code = re.sub(r'\s+', '', ag.cell(r, 3).value).rstrip('.')
    vals = [ag.cell(r, c).value for c in (8, 9, 10)]
    nn = [x for x in vals if x is not None]
    ag_scores.setdefault(ag_alias[a], {})[code] = (max(nn) if nn else None, len(nn) > 1)

perf = wb['Performance Unidades 2023']
perf_alias = {7: 'UBS Jardim Boer', 8: 'UBS Jardim Ipiranga', 9: 'UBS Jardim São Paulo', 10: 'UBS Dona Rosa',
              11: 'UBS Guanabara', 12: 'UBS Mathiensen', 15: 'UBS São Domingos', 17: 'UBS Cariobinha',
              18: 'UBS Vila Gallo', 19: 'ESF Alvorada', 23: 'ESF Mário Covas', 25: 'ESF São José',
              27: 'UAD - Unidade de Atendimento Domiciliar', 29: 'CAPS AD Vida Nova (Álcool e Drogas)',
              32: 'Núcleo de Especialidades de Americana', 33: 'UPA São José', 34: 'PA Antonio Zanaga',
              35: 'UNACON - Unidade de Oncologia', 38: 'CCZ - Centro de Controle de Zoonoses',
              39: 'CTA - Centro de Testagem e Aconselhamento'}
SKIP = {'-', '', 'Não avaliada neste ano', '0'}
visitas = []  # dicts

def add_visit(unid, ciclo, data, situacao, origem, itens_map, obs=None, aprox=False, coments=None):
    visitas.append(dict(id=len(visitas) + 1, unidade=unid, ciclo=ciclo, data=data, aprox=aprox,
                        situacao=situacao, origem=origem, itens=itens_map, obs=obs, coments=coments or {}))

datas_override_2023 = {
    9: (dt.date(2023, 11, 1), True, 'Data registrada como "nov/23??" na Performance 2023; dia aproximado.'),
    35: (dt.date(2023, 12, 8), False, 'Performance 2023 registrava "agosto??"; adotada a data 08/12/2023 da aba UNACON, que traz as mesmas notas.'),
    29: (dt.date(2023, 11, 27), False, 'Performance 2023 registrava 21/11/2023; adotada a data 27/11/2023 da aba da unidade.'),
}
for r, unid in perf_alias.items():
    obs = []
    data = perf.cell(r, 17).value
    aprox = False
    if r in datas_override_2023:
        data, aprox, o = datas_override_2023[r]; obs.append(o)
    if isinstance(data, dt.datetime): data = data.date()
    coments = {}
    for d, col in zip((1, 2, 3, 4), (5, 8, 11, 14)):
        t = perf.cell(r, col).value
        if t and str(t).strip() not in SKIP: coments[d] = str(t).strip()
    if unid in ag_scores:
        m = {}
        for code in ITEM_CODES:
            v, dup = ag_scores[unid].get(code, (None, False))
            m[code] = dict(p=v, na=False)
            if dup: obs.append(f'Item {code} com duas marcações na Avaliação Global (0 e 1); adotado 1, conforme Performance 2023.')
        if unid == 'UPA São José':
            obs.append('Item III.1 = 0 na Avaliação Global, mas a Performance 2023 registrava 1 ponto (50%) digitado manualmente; mantido o valor do item.')
        if unid == 'UBS Jardim São Paulo':
            obs.append('Pontuação recuperada da Avaliação Global; na Performance 2023 a linha apresentava erro de referência (fórmula quebrada).')
        if unid == 'CCZ - Centro de Controle de Zoonoses' and 4 in coments:
            obs.append('Comentário de percepção do usuário idêntico ao da UNACON na Performance 2023 (provável cópia).')
    else:  # ESF Sao Jose: so infraestrutura fisica (14/14 => todos os itens com 2)
        m = {c: dict(p=(2 if c.startswith('I.') else None), na=False) for c in ITEM_CODES}
        obs.append('Apenas a dimensão Infraestrutura Física foi avaliada (14 de 14 pontos); itens reconstruídos com nota máxima.')
    nulos = sum(1 for x in m.values() if x['p'] is None)
    situacao = 'incompleta' if nulos else 'concluida'
    if nulos: obs.append(f'{nulos} item(ns) sem pontuação na origem.')
    add_visit(unid, '2023', data, situacao, f'Planilha: Performance Unidades 2023 (linha {r}) + Avaliação Global',
              m, ' '.join(obs) or None, aprox, coments)

# ------------------------------------------------------------------ visitas 2026 (abas por unidade)
NAO_IMPORTAR = {
    'ESF Zanaga II': 'Pontuação e data (29/06/2023) idênticas às da UAD em 2023: aba copiada sem preenchimento.',
    'ADAM': 'Pontuação e data (29/06/2023) idênticas às da UAD em 2023: aba copiada sem preenchimento.',
    'AEPHIVA': 'Pontuação e data (29/06/2023) idênticas às da UAD em 2023: aba copiada sem preenchimento.',
    'Clínica Jones': 'Pontuação e data (29/06/2023) idênticas às da UAD em 2023: aba copiada sem preenchimento.',
    'UBS Dona Rosa': 'Pontuação idêntica à da UBS Parque Gramado (14/05/2026) com data de 30/10/2023, que não corresponde à visita real de 2023 da unidade.',
    'UBS LIBERDADE': 'Cópia da aba UBS Dona Rosa (mesmo cabeçalho, data e pontuação).',
    'APAE': 'Sem data de visita e com todos os itens em nota máxima: aparenta modelo não preenchido.',
    'UVISA': 'Data 29/06/2023 coincide com a visita da UAD; pontuação não consta da Avaliação Global 2023. Requer confirmação do conselho.',
    'UNACON': 'Notas idênticas à visita de 2023 (já importada pela Avaliação Global).',
    'CTA': 'Notas idênticas à visita de 2023 (já importada pela Avaliação Global).',
    'CAPS Alcool e Drogas': 'Notas idênticas à visita de 2023 (já importada pela Avaliação Global).',
}
NA_ITENS = {'UAD': ['IV.1', 'IV.2'], 'SAE DST': ['IV.1', 'IV.2'], 'Clínica Mod.  ZANAGA': ['IV.2']}
excluidas = []
abas_info = []
for i in range(43):
    ws = wb.worksheets[i]
    aba = ws.title
    start = 22 if ws.cell(22, 1).value == 'Dimensões avaliadas' else 21
    a18 = ws['A18'].value or ''
    mm = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', a18)
    data = dt.date(int(mm.group(3)), int(mm.group(2)), int(mm.group(1))) if mm else None
    m = {}
    for k, r in enumerate(range(4, 18)):
        vals = [ws.cell(r, c).value for c in (7, 8, 9)]
        num = [v for v in vals if isinstance(v, (int, float))]
        txt = [v for v in vals if isinstance(v, str)]
        m[ITEM_CODES[k]] = dict(p=(num[0] if num else None), na=bool(txt))
    for c in NA_ITENS.get(aba, []):
        m[c] = dict(p=None, na=True)
    abas_info.append(dict(aba=aba, cabecalho=ws.cell(start, 5).value, data=data,
                          vetor=''.join('N' if m[c]['na'] else str(m[c]['p']) for c in ITEM_CODES)))
    if aba in NAO_IMPORTAR:
        excluidas.append(dict(aba=aba, unidade=ABA2U[aba], data=data, motivo=NAO_IMPORTAR[aba])); continue
    obs = []
    hdr = str(ws.cell(start, 5).value).strip()
    if aba == 'UAD':
        obs.append('Itens IV.1 e IV.2 estavam com 2 pontos, mas o quadro-resumo marcava a dimensão como "não aplicável"; importados como não se aplica.')
    if aba == 'SAE DST':
        obs.append('Percepção do usuário não se aplica em respeito ao sigilo.')
    if aba == 'Clínica Mod.  ZANAGA':
        obs.append('Item IV.2 marcado como não se aplica.')
    if hdr in ('UBS D. ROSA',) or (aba == 'ESF São José' and 'VITO' in hdr):
        obs.append(f'Cabeçalho do quadro-resumo da aba indicava "{hdr}" (erro de cópia).')
    add_visit(ABA2U[aba], '2026', data, 'concluida', f'Planilha: aba "{aba.strip()}"', m, ' '.join(obs) or None)

# ------------------------------------------------------------------ SQL
ref = ['-- Dados de referencia: instrumento, faixas, tipos e ciclos', 'USE convisita;', 'SET NAMES utf8mb4;', '']
for d in dims:
    ref.append(f"INSERT INTO dimensao (id,codigo,nome,nome_curto,peso,ordem) VALUES ({d[0]},{q(d[1])},{q(d[2])},{q(d[3])},{d[4]},{d[5]});")
for it in itens:
    ref.append(f"INSERT INTO item_indicador (id,dimensao_id,codigo,nome,meta,fonte,periodicidade,pontuacao_max,ordem) VALUES "
               f"({it['id']},{it['dim']},{q(it['codigo'])},{q(it['nome'])},{q(it['meta'])},{q(it['fonte'])},{q(it['periodicidade'])},2,{it['ordem']});")
ref += [
    "INSERT INTO faixa_classificacao VALUES (1,'vermelho','Crítico (abaixo de 40%)',0.0000,0.4000);",
    "INSERT INTO faixa_classificacao VALUES (2,'amarelo','Atenção (40% a 79,9%)',0.4000,0.8000);",
    "INSERT INTO faixa_classificacao VALUES (3,'verde','Adequado (80% ou mais)',0.8000,1.0001);",
]
for n, i in T.items():
    ref.append(f"INSERT INTO tipo_unidade (id,nome) VALUES ({i},{q(n)});")
ref += ["INSERT INTO ciclo (id,nome,data_inicio,data_fim) VALUES (1,'2023','2023-01-01','2023-12-31');",
        "INSERT INTO ciclo (id,nome,data_inicio,data_fim) VALUES (2,'2026','2026-01-01','2026-12-31');",
        "-- usuario inicial: admin / comsaude (TROCAR APOS O PRIMEIRO ACESSO)",
        "INSERT INTO usuario (nome,login,senha_hash,perfil) VALUES ('Administrador','admin','$HASH$','admin');"]
import subprocess
h = subprocess.check_output(['php', '-r', "echo password_hash('comsaude', PASSWORD_DEFAULT);"]).decode()
ref[-1] = ref[-1].replace('$HASH$', h)
open(SQLDIR + '02_referencia.sql', 'w').write('\n'.join(ref) + '\n')

CIC = {'2023': 1, '2026': 2}
IT = {it['codigo']: it['id'] for it in itens}
car = ['-- Carga historica extraida de "Pontuacao visitas 2026atualizada.xlsx"', 'USE convisita;', 'SET NAMES utf8mb4;', '']
for n, t, a, o in unidades:
    car.append(f"INSERT INTO unidade (id,nome,nome_planilha,tipo_id,situacao,periodicidade_meses,observacao) VALUES "
               f"({U[n]},{q(n)},{q(a.strip() if a else None)},{T[t]},'ativa',6,{q(o)});")
car.append('')
for v in visitas:
    car.append(f"INSERT INTO visita (id,uuid,criado_em,unidade_id,ciclo_id,data_visita,data_aproximada,tipo,situacao,origem,observacao_importacao) VALUES "
               f"({v['id']},'hist-{v['id']:04d}','2026-10-03 00:00:00',{U[v['unidade']]},{CIC[v['ciclo']]},{q(v['data'])},{1 if v['aprox'] else 0},'ordinaria',{q(v['situacao'])},{q(v['origem'])},{q(v['obs'])});")
    rows = []
    for c in ITEM_CODES:
        x = v['itens'][c]
        rows.append(f"({v['id']},{IT[c]},{q(x['p'])},{1 if x['na'] else 0})")
    car.append("INSERT INTO avaliacao_item (visita_id,item_id,pontuacao,nao_se_aplica) VALUES " + ','.join(rows) + ';')
    for d, t in v['coments'].items():
        car.append(f"INSERT INTO comentario_dimensao (visita_id,dimensao_id,texto) VALUES ({v['id']},{d},{q(t)});")
open(SQLDIR + '03_carga_historica.sql', 'w').write('\n'.join(car) + '\n')

json.dump(dict(dims=dims, itens=itens, tipos=tipos, unidades=unidades, visitas=visitas, excluidas=excluidas,
               abas=abas_info), open('carga.json', 'w'), ensure_ascii=False, default=str, indent=1)
print(len(unidades), 'unidades', len(visitas), 'visitas', len(excluidas), 'excluidas')
for v in visitas: print(v['ciclo'], v['data'], v['situacao'], v['unidade'], '|', (v['obs'] or '')[:90])
